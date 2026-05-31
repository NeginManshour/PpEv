#!/usr/bin/env python
# coding: utf-8

# In[ ]:


import json

# Load configuration and define a global variable
def load_config():
    global config
    with open('config_Pyrosetta.json') as config_file:
        config = json.load(config_file)

# Call the function to load the configuration
load_config()


# In[ ]:


import os
import json
import pandas as pd

# Load configuration and define a global variable
def load_config():
    global config
    with open('config_Pyrosetta.json') as config_file:
        config = json.load(config_file)

# Call the function to load the configuration
load_config()

def convert_tsv_to_excel(source_directory, target_directory, subdirectory_name):
    """
    Converts all TSV files found in the source_directory into Excel files
    and saves them into a new subdirectory inside the target_directory.

    Args:
    - source_directory (str): The path to the directory containing TSV files.
    - target_directory (str): The path to the directory where the new subdirectory will be created.
    - subdirectory_name (str): The name of the new subdirectory where Excel files will be saved.
    """
    # Create the new subdirectory path
    new_directory_path = os.path.join(target_directory, subdirectory_name)

    # Ensure the new subdirectory exists
    if not os.path.exists(new_directory_path):
        os.makedirs(new_directory_path)

    # Iterate over all files in the source directory
    for filename in os.listdir(source_directory):
        if filename.endswith('.tsv'):
            # Construct the full file paths
            source_file_path = os.path.join(source_directory, filename)
            # Change the file extension from .tsv to .xlsx for the output file
            target_file_path = os.path.join(new_directory_path, filename.replace('.tsv', '.xlsx'))

            # Load the TSV file
            df = pd.read_csv(source_file_path, sep='\t')
            # Save the dataframe to an Excel file
            df.to_excel(target_file_path, index=False)

            #print(f'Converted {filename} to Excel and saved as {os.path.basename(target_file_path)}')

# Example usage
source_directory = config['Pyrosetta_directory']
target_directory = config['output_directories']
subdirectory_name = 'Pyros_excel_TB'

convert_tsv_to_excel(source_directory, target_directory, subdirectory_name)


# In[ ]:


import os
import pandas as pd

# Directory containing the Excel files
#directory_path = '/Users/neginmanshour/Desktop/Protein_Peptide_Evaluation/Pyrosetta/Pyros_ex_TF'
directory_path = f"{config['output_directories']}/Pyros_excel_TB"

# List all Excel files in the directory
excel_files = [file for file in os.listdir(directory_path) if file.endswith('.xlsx')]

for file_name in excel_files:
    file_path = os.path.join(directory_path, file_name)

    try:
        # Load the Excel file
        df = pd.read_excel(file_path)

        # Check if the DataFrame has at least 4 columns
        if df.shape[1] < 4:
            print(f"File {file_name} does not have enough columns.")
            continue

        # Rename the columns
        df.columns = ['Pyrosetta_rank', 'alphafold_rank', 'structure_name', 'score']

        # Save the modified DataFrame back to the same file
        df.to_excel(file_path, index=False)
        #print(f"Updated columns for file {file_name}")

    except Exception as e:
        print(f"An error occurred with file {file_name}: {e}")

print("Column renaming completed for all files.")



# In[ ]:


import os
import shutil
import pandas as pd
from openpyxl import load_workbook

def extract_pdb_id_from_dockq(filename):
    """Extracts the PDB ID from a DockQ filename."""
    parts = filename.split('_')
    return parts[0]  # PDB ID is the first part

def extract_pdb_id_from_pyrosetta(filename):
    """Extracts the PDB ID from a PyRosetta filename."""
    parts = filename.split('_')
    return parts[2]  # PDB ID is after 'pyrosetta_scoring'

def copy_pyrosetta_sheets_to_dockq(dockq_dir, pyrosetta_dir, target_directory, subdirectory_name):
    """
    Copies PyRosetta sheets to corresponding DockQ files and saves the results in a new subdirectory inside target_directory.

    Args:
    - dockq_dir (str): Directory containing DockQ files.
    - pyrosetta_dir (str): Directory containing PyRosetta files.
    - target_directory (str): Directory where the new subdirectory will be created.
    - subdirectory_name (str): Name of the new subdirectory for saving the results.
    """
    # Create the new subdirectory path
    output_dir = os.path.join(target_directory, subdirectory_name)

    # Ensure the new subdirectory exists
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)

    dockq_files = os.listdir(dockq_dir)
    pyrosetta_files = os.listdir(pyrosetta_dir)

    # Map PyRosetta files to their PDB IDs
    pyrosetta_map = {extract_pdb_id_from_pyrosetta(f).upper(): f for f in pyrosetta_files if 'pyrosetta_scoring' in f}

    for dockq_file in dockq_files:
        if dockq_file.endswith('_TB.xlsx') and 'DockQ_data' in dockq_file:
            pdb_id = extract_pdb_id_from_dockq(dockq_file).upper()
            if pdb_id in pyrosetta_map:
                dockq_path = os.path.join(dockq_dir, dockq_file)
                pyrosetta_path = os.path.join(pyrosetta_dir, pyrosetta_map[pdb_id])
                output_path = os.path.join(output_dir, dockq_file)

                # Copy the DockQ file to the output directory if it's not already there
                if not os.path.exists(output_path):
                    shutil.copyfile(dockq_path, output_path)

                print(f"Processing {dockq_file} and {pyrosetta_map[pdb_id]} for PDB ID {pdb_id}")

                pyrosetta_df = pd.read_excel(pyrosetta_path, sheet_name='Sheet1')

                with pd.ExcelWriter(output_path, engine='openpyxl', mode='a') as writer:
                    book = load_workbook(output_path)
                    writer.book = book
                    if 'Pyrosetta' in book.sheetnames:
                        std = book['Pyrosetta']
                        book.remove(std)
                    pyrosetta_df.to_excel(writer, sheet_name='Pyrosetta', index=False)

                    print(f'Updated {dockq_file} with PyRosetta data for PDB ID {pdb_id}')
            else:
                print(f"No matching PyRosetta file found for {dockq_file}")

# Replace these paths with the actual paths to your directories
dockq_dir = config['TB_uni']
pyrosetta_dir = f"{config['output_directories']}/Pyros_excel_TB"
target_directory = config['output_directories']
subdirectory_name = 'Pyrosetta_DockQ_TB'

copy_pyrosetta_sheets_to_dockq(dockq_dir, pyrosetta_dir, target_directory, subdirectory_name)


# In[ ]:


import os
import shutil
import pandas as pd
from openpyxl import load_workbook
import json


def extract_pdb_id_from_dockq(filename):
    """Extracts the PDB ID from a DockQ filename."""
    parts = filename.split('_')
    return parts[0]  # PDB ID is the first part

def extract_pdb_id_from_pyrosetta(filename):
    """Extracts the PDB ID from a PyRosetta filename."""
    parts = filename.split('_')
    return parts[2]  # PDB ID is after 'pyrosetta_scoring'

def copy_pyrosetta_sheets_to_dockq(dockq_dir, pyrosetta_dir, target_directory, subdirectory_name):
    """
    Copies PyRosetta sheets to corresponding DockQ files and saves the results in a new subdirectory inside target_directory.

    Args:
    - dockq_dir (str): Directory containing DockQ files.
    - pyrosetta_dir (str): Directory containing PyRosetta files.
    - target_directory (str): Directory where the new subdirectory will be created.
    - subdirectory_name (str): Name of the new subdirectory for saving the results.
    """
    # Create the new subdirectory path
    output_dir = os.path.join(target_directory, subdirectory_name)

    # Ensure the new subdirectory exists
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)

    dockq_files = os.listdir(dockq_dir)
    pyrosetta_files = os.listdir(pyrosetta_dir)

    # Map PyRosetta files to their PDB IDs
    pyrosetta_map = {extract_pdb_id_from_pyrosetta(f).upper(): f for f in pyrosetta_files if 'pyrosetta_scoring' in f}

    for dockq_file in dockq_files:
        if dockq_file.endswith('_TB.xlsx') and 'DockQ_data' in dockq_file:
            pdb_id = extract_pdb_id_from_dockq(dockq_file).upper()
            if pdb_id in pyrosetta_map:
                dockq_path = os.path.join(dockq_dir, dockq_file)
                pyrosetta_path = os.path.join(pyrosetta_dir, pyrosetta_map[pdb_id])
                output_path = os.path.join(output_dir, dockq_file)

                # Copy the DockQ file to the output directory if it's not already there
                if not os.path.exists(output_path):
                    shutil.copyfile(dockq_path, output_path)

                print(f"Processing {dockq_file} and {pyrosetta_map[pdb_id]} for PDB ID {pdb_id}")

                pyrosetta_df = pd.read_excel(pyrosetta_path, sheet_name='Sheet1')

                # Load the existing workbook
                book = load_workbook(output_path)

                # Remove the existing sheet if it exists
                if 'Pyrosetta' in book.sheetnames:
                    std = book['Pyrosetta']
                    book.remove(std)

                # Save the updated workbook
                book.save(output_path)

                # Add the new sheet to the workbook
                with pd.ExcelWriter(output_path, engine='openpyxl', mode='a') as writer:
                    pyrosetta_df.to_excel(writer, sheet_name='Pyrosetta', index=False)

                print(f'Updated {dockq_file} with PyRosetta data for PDB ID {pdb_id}')
            else:
                print(f"No matching PyRosetta file found for {dockq_file}")

# Use the paths from the already loaded config dictionary
dockq_dir = config['TB_uni']
pyrosetta_dir = f"{config['output_directories']}/Pyros_excel_TB"
target_directory = config['output_directories']
subdirectory_name = 'Pyrosetta_DockQ_TB'

copy_pyrosetta_sheets_to_dockq(dockq_dir, pyrosetta_dir, target_directory, subdirectory_name)


# In[ ]:


import os
import pandas as pd

def calculate_spearman_for_directory(directory_path, output_file):
    # List to store results
    results = []

    # Iterate through all files in the directory
    for filename in os.listdir(directory_path):
        if filename.endswith(".xlsx"):
            file_path = os.path.join(directory_path, filename)

            # Load the Excel file
            excel_data = pd.ExcelFile(file_path)

            # Check if the expected sheets exist
            if 'Sheet' in excel_data.sheet_names and 'Pyrosetta' in excel_data.sheet_names:
                # Load the data from each sheet
                sheet_data = pd.read_excel(file_path, sheet_name='Sheet')
                pyrosetta_data = pd.read_excel(file_path, sheet_name='Pyrosetta')

                # Extract relevant columns and rename them for clarity
                sheet_relevant = sheet_data[['File Name', 'DockQ Rank']].rename(columns={'File Name': 'structure_name', 'DockQ Rank': 'dock_rank'})
                pyrosetta_relevant = pyrosetta_data[['structure_name', 'Pyrosetta_rank']]

                # Merge the dataframes based on 'structure_name'
                merged_data = pd.merge(sheet_relevant, pyrosetta_relevant, on='structure_name')

                # Calculate the Spearman correlation
                spearman_corr = merged_data['dock_rank'].corr(merged_data['Pyrosetta_rank'], method='spearman')

                # Append the result to the list
                results.append({'File Name': filename, 'Spearman Correlation': spearman_corr})

    # Create a DataFrame from the results
    results_df = pd.DataFrame(results)

    # Save the results to an Excel file
    results_df.to_excel(output_file, index=False)

# Usage
directory_path = config['output_directories'] + "/Pyrosetta_DockQ_TB"
output_directory = config['Spearman_Correlation_directory_TB']
output_file_path = os.path.join(output_directory, "spearman_correlation_Pyrosetta_TB.xlsx")

calculate_spearman_for_directory(directory_path, output_file_path)


# In[ ]:


import pandas as pd

# Load the Excel file
# Make sure to include the file extension, such as '.xlsx', in the file path
file_path = f"{config['Spearman_Correlation_directory_TB']}/spearman_correlation_Pyrosetta_TB.xlsx" # Adjusted file path with extension
try:
    df = pd.read_excel(file_path)

    # Calculate the percentages of positive and negative Spearman correlation values
    positive_percentage = (df['Spearman Correlation'] > 0).mean() * 100
    negative_percentage = (df['Spearman Correlation'] < 0).mean() * 100

    # Calculate the average of Spearman Correlation values
    average_correlation = df['Spearman Correlation'].mean()

    # Print the results
    print(f"Percentage of positive correlations: {positive_percentage:.2f}%")
    print(f"Percentage of negative correlations: {negative_percentage:.2f}%")
    print(f"Average Spearman Correlation: {average_correlation:.2f}")
except FileNotFoundError:
    print(f"File not found. Please check the file path: {file_path}")


# In[ ]:


import os
import pandas as pd

# Directory containing the Excel files
directory_path = f"{config['output_directories']}/Pyrosetta_DockQ_TB"

# Initialize a list to store the results
results = []

# List all Excel files in the directory
excel_files = [file for file in os.listdir(directory_path) if file.endswith('.xlsx')]

for file_name in excel_files:
    file_path = os.path.join(directory_path, file_name)
    xls = pd.ExcelFile(file_path)

    # Extract pdb_id from file name
    pdb_id = file_name.split('_')[0]

    # Load data from "Pyrosetta" sheet
    Pyrosetta_df = pd.read_excel(file_path, sheet_name='Pyrosetta')
    structure_name_for_rank_zero = Pyrosetta_df[Pyrosetta_df['Pyrosetta_rank'] == 0]['structure_name'].iloc[0]

    # Load data from "Sheet" sheet
    sheet_df = pd.read_excel(file_path, sheet_name='Sheet')

    # Highest DockQ score
    highest_dockq_score = sheet_df['DockQ'].max()

    # DockQ score for the structure from Pyrosetta
    highest_Pyrosetta_dockq_score = sheet_df[sheet_df['File Name'].str.contains(structure_name_for_rank_zero)]['DockQ'].max()

    # Calculate Loss
    score_Loss = highest_dockq_score - highest_Pyrosetta_dockq_score

    # Append results
    results.append([pdb_id, highest_dockq_score, highest_Pyrosetta_dockq_score, score_Loss])

# Create a DataFrame from the results
results_df = pd.DataFrame(results, columns=['File Name', 'DockQ', 'Pyrosetta ranked', 'Loss'])

# Save the results into a new Excel file
output_directory = config["DockQ_Loss_directory_TB"]
output_path = os.path.join(output_directory, 'Pyrosetta_TB.xlsx')
results_df.to_excel(output_path, index=False)

print(f"Results have been saved to {output_path}")


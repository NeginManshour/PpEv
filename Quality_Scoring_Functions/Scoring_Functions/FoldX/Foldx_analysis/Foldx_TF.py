#!/usr/bin/env python
# coding: utf-8

# In[ ]:


import json

# Load configuration and define a global variable
def load_config():
    global config
    with open('config_FoldX.json') as config_file:
        config = json.load(config_file)

# Call the function to load the configuration
load_config()


# In[ ]:


import os
import shutil
import pandas as pd
from openpyxl import load_workbook
import json

# Load configuration and define a global variable
def load_config():
    global config
    with open('config_FoldX.json') as config_file:
        config = json.load(config_file)

# Call the function to load the configuration
load_config()

def extract_pdb_id_from_dockq(filename):
    """Extracts the PDB ID from a DockQ filename."""
    parts = filename.split('_')
    return parts[0]  # PDB ID is the first part

def extract_pdb_id_from_foldx(filename):
    """Extracts the PDB ID from a foldx filename."""
    parts = filename.split('_')
    return parts[2]  # PDB ID is after 'foldx_scoring'

def copy_foldx_sheets_to_dockq(dockq_dir, foldx_dir, output_dir):
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)

    dockq_files = os.listdir(dockq_dir)
    foldx_files = os.listdir(foldx_dir)

    # Map foldx files to their PDB IDs
    foldx_map = {extract_pdb_id_from_foldx(f).upper(): f for f in foldx_files if 'binding_data' in f}

    for dockq_file in dockq_files:
        if dockq_file.endswith('_TF.xlsx') and 'DockQ_data' in dockq_file:
            pdb_id = extract_pdb_id_from_dockq(dockq_file).upper()
            if pdb_id in foldx_map:
                dockq_path = os.path.join(dockq_dir, dockq_file)
                foldx_path = os.path.join(foldx_dir, foldx_map[pdb_id])
                output_path = os.path.join(output_dir, dockq_file)

                # Copy the DockQ file to the output directory if it's not already there
                if not os.path.exists(output_path):
                    shutil.copyfile(dockq_path, output_path)

                print(f"Processing {dockq_file} and {foldx_map[pdb_id]} for PDB ID {pdb_id}")

                foldx_df = pd.read_excel(foldx_path, sheet_name='Sheet1')

                # Load the existing workbook
                book = load_workbook(output_path)

                # Remove the existing sheet if it exists
                if 'Foldx' in book.sheetnames:
                    std = book['Foldx']
                    book.remove(std)

                # Save the updated workbook
                book.save(output_path)

                # Add the new sheet to the workbook
                with pd.ExcelWriter(output_path, engine='openpyxl', mode='a') as writer:
                    foldx_df.to_excel(writer, sheet_name='Foldx', index=False)

                print(f'Updated {dockq_file} with Foldx data for PDB ID {pdb_id}')
            else:
                print(f"No matching Foldx file found for {dockq_file}")

# Use paths from the configuration file
dockq_dir = config["TF_uni"]
foldx_dir = config["source_directory"]
output_dir = os.path.join(config["output_directory"], "foldx_TF_Dock")

copy_foldx_sheets_to_dockq(dockq_dir, foldx_dir, output_dir)


# In[ ]:


import pandas as pd
import os
import re

# Define the directories

input_directory = os.path.join(config['output_directory'], 'foldx_TF_Dock')
output_directory = os.path.join(config['output_directory'], 'alphafold_rank_Foldx_TF')


# Ensure the output directory exists
if not os.path.exists(output_directory):
    os.makedirs(output_directory)

# Function to extract the rank number
def extract_rank(s):
    match = re.search(r'\d+', s)
    return int(match.group()) if match else None

# Process each Excel file in the directory
for filename in os.listdir(input_directory):
    if filename.endswith('.xlsx') or filename.endswith('.xls'):
        file_path = os.path.join(input_directory, filename)
        output_file_path = os.path.join(output_directory, filename)

        # Load the workbook
        xls = pd.ExcelFile(file_path)

        # Initialize a dict to hold dataframes
        dfs = {}
        for sheet_name in xls.sheet_names:
            # Read each sheet into a dataframe
            df = xls.parse(sheet_name)

            # Modify the "Foldx" sheet
            if sheet_name == 'Foldx':
                # Extract "alphafold rank" and reorder columns
                df['alphafold rank'] = df['Unnamed: 0'].apply(extract_rank)
                cols = ['alphafold rank'] + [col for col in df if col != 'alphafold rank']
                df = df[cols]

                # Delete rows where "Unnamed: 0" contains "native"
                # Adjust the column name if the information is located in a different column
                df = df[~df['Unnamed: 0'].str.contains("native", case=False, na=False)]

            # Store the modified (or unmodified) dataframe
            dfs[sheet_name] = df

        # Write all dataframes to a new Excel file
        with pd.ExcelWriter(output_file_path, engine='openpyxl') as writer:
            for sheet_name, df in dfs.items():
                df.to_excel(writer, sheet_name=sheet_name, index=False)

print("All files have been processed and saved to the output directory.")




# In[ ]:


import pandas as pd
import os
import re

def process_excel_files(input_directory, output_directory):
    # Ensure output directory exists
    if not os.path.exists(output_directory):
        os.makedirs(output_directory)

    # Iterate through all Excel files in the input directory
    for filename in os.listdir(input_directory):
        if filename.endswith('.xlsx'):
            file_path = os.path.join(input_directory, filename)
            output_file_path = os.path.join(output_directory, filename)

            # Load the Excel file
            xls = pd.ExcelFile(file_path)

            # Create a Pandas Excel writer using XlsxWriter as the engine.
            with pd.ExcelWriter(output_file_path, engine='xlsxwriter') as writer:
                # Process each sheet
                for sheet_name in xls.sheet_names:
                    df = pd.read_excel(file_path, sheet_name=sheet_name)

                    # Specific operations for the 'Foldx' sheet
                    if sheet_name == 'Foldx':
                        # Extracting alphafold rank from 'Unnamed: 0' and adding it as a new column if not exists
                        if 'alphafold rank' not in df.columns:
                            df['alphafold rank'] = df['Unnamed: 0'].apply(lambda x: int(re.search(r'\d+', x).group()) if re.search(r'\d+', x) else 0)

                        # Delete rows related to 'native'
                        df = df[df['Unnamed: 0'].str.contains('native') == False]

                        # Rank data based on Stability column
                        df = df.sort_values(by='Stability')
                        df['stability_rank'] = range(len(df))

                    # Write each DataFrame to a specific sheet
                    df.to_excel(writer, sheet_name=sheet_name, index=False)

# Example usage
input_directory = os.path.join(config['output_directory'], 'alphafold_rank_Foldx_TF')
output_directory = os.path.join(config['output_directory'], 'stability_TF')
process_excel_files(input_directory, output_directory)


# In[ ]:


import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows

# Define the source and output directories

source_directory = os.path.join(config['output_directory'], 'alphafold_rank_Foldx_TF')
output_directory = os.path.join(config['output_directory'], 'interaction_TF')

if not os.path.exists(output_directory):
    os.makedirs(output_directory)

# Iterate over all Excel files in the source directory
for filename in os.listdir(source_directory):
    if filename.endswith('.xlsx'):
        file_path = os.path.join(source_directory, filename)

        # Load the workbook and sheet
        workbook = load_workbook(file_path)
        if 'Foldx' in workbook.sheetnames:
            ws = workbook['Foldx']
            data = ws.values
            columns = next(data)[0:]  # First row for column names
            df = pd.DataFrame(data, columns=columns)
            df = df[1:]  # Remove the header row from the data

            # Convert 'interaction_energy' to numeric and sort
            df['interaction_energy'] = pd.to_numeric(df['interaction_energy'], errors='coerce')
            df_sorted = df.sort_values(by='interaction_energy').reset_index(drop=True)
            df_sorted['interaction_energy_rank'] = range(len(df_sorted))

            # Clear existing data in the sheet
            for row in ws['A2:Z' + str(ws.max_row)]:
                for cell in row:
                    cell.value = None

            # Write the updated DataFrame back to the Excel sheet, including header
            for r_idx, row in enumerate(dataframe_to_rows(df_sorted, index=False, header=True), start=1):
                for c_idx, val in enumerate(row, start=1):
                    ws.cell(row=r_idx, column=c_idx, value=val)

            # Save the workbook to a new file in the output directory
            output_file_path = os.path.join(output_directory, filename)
            workbook.save(output_file_path)

print("Processing complete. All files have been updated.")


# In[ ]:


import os
import openpyxl

# Set the directory containing the Excel files
directory = os.path.join(config['output_directory'], 'interaction_TF')

# Iterate through each file in the directory
for filename in os.listdir(directory):
    if filename.endswith('.xlsx'):
        # Construct the full file path
        file_path = os.path.join(directory, filename)
        # Load the workbook
        workbook = openpyxl.load_workbook(file_path)

        # Check if 'Foldx' sheet is in the workbook
        if 'Foldx' in workbook.sheetnames:
            # Rename the sheet
            workbook["Foldx"].title = "Foldx_Int"
            # Save the changes to the same file
            workbook.save(file_path)
            #print(f'Renamed sheet in {filename}')
        else:
            print(f'No sheet named "Foldx" in {filename}')

print("Sheet renaming complete.")


# In[ ]:


import os
import pandas as pd
import matplotlib.pyplot as plt
from scipy.stats import spearmanr
from zipfile import BadZipFile

def process_excel_files(directory_path, output_directory):
    excel_files = [f for f in os.listdir(directory_path) if f.endswith('.xlsx')]
    correlations = []  # To store file names and their Spearman correlations

    for excel_file in excel_files:
        excel_file_path = os.path.join(directory_path, excel_file)

        try:
            xls = pd.ExcelFile(excel_file_path)
        except BadZipFile:
            print(f'Error: The file "{excel_file}" is not a valid Excel file or is corrupted.')
            continue
        except Exception as e:
            print(f'An unexpected error occurred while processing {excel_file}: {e}')
            continue

        if 'Foldx' not in xls.sheet_names:
            print(f'The file "{excel_file}" does not contain a "Foldx" sheet.')
            continue

        dockq_df = pd.read_excel(excel_file_path, sheet_name='Sheet')
        Foldx_df = pd.read_excel(excel_file_path, sheet_name='Foldx')

        Foldx_rank_column = 'stability_rank'
        alpha_to_dockq_map = dockq_df.set_index('AlphaFold Rank')['DockQ Rank'].to_dict()
        Foldx_df['DockQ Rank'] = Foldx_df['alphafold rank'].map(alpha_to_dockq_map)

        filtered_Foldx_df = Foldx_df.dropna(subset=['DockQ Rank'])

        spearman_corr, _ = spearmanr(filtered_Foldx_df[Foldx_rank_column], filtered_Foldx_df['DockQ Rank'])

        plt.figure(figsize=(10, 6))
        plt.scatter(filtered_Foldx_df['DockQ Rank'], filtered_Foldx_df[Foldx_rank_column], alpha=0.6)
        plt.title(f'File: {excel_file}\nSpearman Correlation: {spearman_corr:.2f}')
        plt.xlabel('DockQ Rank')
        plt.ylabel('stability_rank')
        plt.grid(True)
        plt.show()

        # Append file name and Spearman correlation to the list
        correlations.append({'File Name': excel_file, 'Spearman Correlation': spearman_corr})
        print(f'File: {excel_file}\nSpearman Correlation Coefficient: {spearman_corr}\n')

    # Convert list to DataFrame
    correlations_df = pd.DataFrame(correlations)

    # Ensure the output directory exists
    if not os.path.exists(output_directory):
        os.makedirs(output_directory)

    # Path for the new Excel file
    output_file_path = os.path.join(output_directory, 'correlations_stability_TF.xlsx')

    # Save DataFrame to an Excel file
    correlations_df.to_excel(output_file_path, index=False)
    print(f'Summary of Spearman correlations saved to {output_file_path}')

# Example usage
directory_path = os.path.join(config['output_directory'], 'stability_TF')
output_directory = config['Spearman_Correlation_directory_TF']
process_excel_files(directory_path, output_directory)



# In[ ]:


import pandas as pd

def calculate_positive_negative_percentages_and_average(file_path):
    # Load the Excel file
    data = pd.read_excel(file_path)

    # Count the total number of entries
    total_entries = len(data['Spearman Correlation'])

    # Count the number of positive and negative Foldx_Spearman Correlation values
    positive_count = data[data['Spearman Correlation'] > 0].shape[0]
    negative_count = data[data['Spearman Correlation'] < 0].shape[0]

    # Calculate percentages
    positive_percentage = (positive_count / total_entries) * 100
    negative_percentage = (negative_count / total_entries) * 100

    # Calculate the average of Foldx_Spearman Correlation values
    average_correlation = data['Spearman Correlation'].mean()

    return positive_percentage, negative_percentage, average_correlation

# Example usage
#file_path = '/Users/neginmanshour/Desktop/Protein_Peptide_Evaluation/spearman_Correlation/TB_Spearman/correlations_stability_TB.xlsx'  # Replace with your actual file path
output_directory = config['Spearman_Correlation_directory_TF']
file_name = 'correlations_stability_TF.xlsx'
file_path = os.path.join(output_directory, file_name)

positive_percentage, negative_percentage, average_correlation = calculate_positive_negative_percentages_and_average(file_path)
print(f"Positive Foldx_Spearman Correlation: {positive_percentage:.2f}%")
print(f"Negative Foldx_Spearman Correlation: {negative_percentage:.2f}%")
print(f"Average Foldx_Spearman Correlation: {average_correlation:.2f}")


# In[ ]:


import os
import pandas as pd

# Directory containing the Excel files
directory_path = os.path.join(config['output_directory'], 'stability_TF')

# Initialize a list to store the results
results = []

# List all Excel files in the directory
excel_files = [file for file in os.listdir(directory_path) if file.endswith('.xlsx')]

for file_name in excel_files:
    file_path = os.path.join(directory_path, file_name)
    xls = pd.ExcelFile(file_path)

    # Extract pdb_id from file name
    pdb_id = file_name.split('_')[0]

    # Load data from "Foldx" sheet
    foldx_df = pd.read_excel(file_path, sheet_name='Foldx')
    structure_name_for_rank_zero = foldx_df[foldx_df['stability_rank'] == 0]['Unnamed: 0'].iloc[0].replace('_clean.pdb', '')

    # Load data from "Sheet" sheet
    sheet_df = pd.read_excel(file_path, sheet_name='Sheet')

    # Highest DockQ score
    highest_dockq_score = sheet_df['DockQ'].max()

    # DockQ score for the structure from Foldx
    highest_foldx_dockq_score = sheet_df[sheet_df['File Name'].str.contains(structure_name_for_rank_zero)]['DockQ'].max()

    # Calculate Loss
    score_Loss = highest_dockq_score - highest_foldx_dockq_score

    # Append results
    results.append([pdb_id, highest_dockq_score, highest_foldx_dockq_score, score_Loss])

# Create a DataFrame from the results
results_df = pd.DataFrame(results, columns=['File Name', 'DockQ', 'Foldx/Stability ranked', 'Loss'])

# Save the results into a new Excel file

output_path = os.path.join(config['DockQ_Loss_directory_TF'], 'Foldx_stability_TF.xlsx')
results_df.to_excel(output_path, index=False)

print(f"Results have been saved to {output_path}")



# In[ ]:


import os
import pandas as pd
import matplotlib.pyplot as plt
from scipy.stats import spearmanr
from zipfile import BadZipFile

def process_excel_files(directory_path, output_directory):
    excel_files = [f for f in os.listdir(directory_path) if f.endswith('.xlsx')]
    correlations = []  # To store file names and their Spearman correlations

    for excel_file in excel_files:
        excel_file_path = os.path.join(directory_path, excel_file)

        try:
            xls = pd.ExcelFile(excel_file_path)
        except BadZipFile:
            print(f'Error: The file "{excel_file}" is not a valid Excel file or is corrupted.')
            continue
        except Exception as e:
            print(f'An unexpected error occurred while processing {excel_file}: {e}')
            continue

        if 'Foldx_Int' not in xls.sheet_names:
            print(f'The file "{excel_file}" does not contain a "Foldx_Int" sheet.')
            continue

        dockq_df = pd.read_excel(excel_file_path, sheet_name='Sheet')
        Foldx_df = pd.read_excel(excel_file_path, sheet_name='Foldx_Int')

        Foldx_rank_column = 'interaction_energy_rank'
        alpha_to_dockq_map = dockq_df.set_index('AlphaFold Rank')['DockQ Rank'].to_dict()
        Foldx_df['DockQ Rank'] = Foldx_df['alphafold rank'].map(alpha_to_dockq_map)

        filtered_Foldx_df = Foldx_df.dropna(subset=['DockQ Rank'])

        spearman_corr, _ = spearmanr(filtered_Foldx_df[Foldx_rank_column], filtered_Foldx_df['DockQ Rank'])

        plt.figure(figsize=(10, 6))
        plt.scatter(filtered_Foldx_df['DockQ Rank'], filtered_Foldx_df[Foldx_rank_column], alpha=0.6)
        plt.title(f'File: {excel_file}\nSpearman Correlation: {spearman_corr:.2f}')
        plt.xlabel('DockQ Rank')
        plt.ylabel('interaction_energy_rank')
        plt.grid(True)
        plt.show()

        # Append file name and Spearman correlation to the list
        correlations.append({'File Name': excel_file, 'Spearman Correlation': spearman_corr})
        print(f'File: {excel_file}\nSpearman Correlation Coefficient: {spearman_corr}\n')

    # Convert list to DataFrame
    correlations_df = pd.DataFrame(correlations)

    # Ensure the output directory exists
    if not os.path.exists(output_directory):
        os.makedirs(output_directory)

    # Path for the new Excel file
    output_file_path = os.path.join(output_directory, 'correlations_Interaction_TF.xlsx')

    # Save DataFrame to an Excel file
    correlations_df.to_excel(output_file_path, index=False)
    print(f'Summary of Spearman correlations saved to {output_file_path}')

# Example usage
directory_path = os.path.join(config['output_directory'], 'interaction_TF')
output_directory = config['Spearman_Correlation_directory_TF']
process_excel_files(directory_path, output_directory)



# In[ ]:


import pandas as pd

def calculate_positive_negative_percentages_and_average(file_path):
    # Load the Excel file
    data = pd.read_excel(file_path)

    # Count the total number of entries
    total_entries = len(data['Spearman Correlation'])

    # Count the number of positive and negative Foldx_Spearman Correlation values
    positive_count = data[data['Spearman Correlation'] > 0].shape[0]
    negative_count = data[data['Spearman Correlation'] < 0].shape[0]

    # Calculate percentages
    positive_percentage = (positive_count / total_entries) * 100
    negative_percentage = (negative_count / total_entries) * 100

    # Calculate the average of Foldx_Spearman Correlation values
    average_correlation = data['Spearman Correlation'].mean()

    return positive_percentage, negative_percentage, average_correlation

# Example usage
#file_path = '/Users/neginmanshour/Desktop/Protein_Peptide_Evaluation/spearman_Correlation/TB_Spearman/correlations_Interaction_TB.xlsx'  # Replace with your actual file path

output_directory = config['Spearman_Correlation_directory_TF']
file_name = 'correlations_Interaction_TF.xlsx'
file_path = os.path.join(output_directory, file_name)

positive_percentage, negative_percentage, average_correlation = calculate_positive_negative_percentages_and_average(file_path)
print(f"Positive Foldx_Spearman Correlation: {positive_percentage:.2f}%")
print(f"Negative Foldx_Spearman Correlation: {negative_percentage:.2f}%")
print(f"Average Foldx_Spearman Correlation: {average_correlation:.2f}")


# In[ ]:


import os
import pandas as pd

# Directory containing the Excel files
#directory_path = '/Users/neginmanshour/Desktop/Protein_Peptide_Evaluation/Foldx/Data/interaction_TB'
directory_path = os.path.join(config['output_directory'], 'interaction_TF')

# Initialize a list to store the results
results = []

# List all Excel files in the directory
excel_files = [file for file in os.listdir(directory_path) if file.endswith('.xlsx')]

for file_name in excel_files:
    file_path = os.path.join(directory_path, file_name)
    xls = pd.ExcelFile(file_path)

    # Extract pdb_id from file name
    pdb_id = file_name.split('_')[0]

    # Load data from "Foldx" sheet
    foldx_df = pd.read_excel(file_path, sheet_name='Foldx_Int')
    structure_name_for_rank_zero = foldx_df[foldx_df['interaction_energy_rank'] == 0]['Unnamed: 0'].iloc[0].replace('_clean.pdb', '')

    # Load data from "Sheet" sheet
    sheet_df = pd.read_excel(file_path, sheet_name='Sheet')

    # Highest DockQ score
    highest_dockq_score = sheet_df['DockQ'].max()

    # DockQ score for the structure from Foldx
    highest_foldx_dockq_score = sheet_df[sheet_df['File Name'].str.contains(structure_name_for_rank_zero)]['DockQ'].max()

    # Calculate Loss
    score_Loss = highest_dockq_score - highest_foldx_dockq_score

    # Append results
    results.append([pdb_id, highest_dockq_score, highest_foldx_dockq_score, score_Loss])

# Create a DataFrame from the results
results_df = pd.DataFrame(results, columns=['File Name', 'DockQ', 'Foldx/Interaction ranked', 'Loss'])

# Save the results into a new Excel file
#output_path = '/Users/neginmanshour/Desktop/Protein_Peptide_Evaluation/Loss/TB_Loss/Foldx_Interaction_TB.xlsx'
output_path = os.path.join(config['DockQ_Loss_directory_TF'], 'Foldx_Interaction_TF.xlsx')
results_df.to_excel(output_path, index=False)

print(f"Results have been saved to {output_path}")


# In[ ]:




